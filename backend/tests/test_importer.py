"""xlsx importer 测试。

使用一份最小化的 xlsx fixture 验证：
- 6 张表都能导入
- AST 缓存正确写入
- 解析失败优雅降级
"""
from __future__ import annotations

import json
import tempfile
from pathlib import Path

import openpyxl

from app.db.models import (
    AdrgRule,
    CcList,
    DrgRule,
    ExclusionTable,
    MdcRule,
    SetMember,
)
from app.db.session import SessionLocal, init_db
from app.engine.parser.ast_nodes import ConstTrueNode, InCheckNode
from app.importers.xlsx_importer import import_rules_xlsx


def _make_minimal_xlsx(path: Path) -> None:
    """构造一份最小 xlsx：6 张表各 2-3 行。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    md = wb.create_sheet("MDC")
    md.append(["标题"])
    md.append(["MDC编码", "MDC名称", "MDC规则", "排序"])
    md.append(["MDCA", "先期分组", "", 1])
    md.append(["MDCB", "神经系统疾病", "ZYZD in DI_B00", 2])

    ad = wb.create_sheet("ADRG")
    ad.append(["标题"])
    ad.append(["ADRG编码", "ADRG名称", "ADRG规则", "所属MDC编码", "MDC内排序"])
    ad.append(["AA1", "心肺移植", "ZYSS in OP_AA1", "MDCA", 1])
    ad.append(["BB1", "兜底组", "1", "MDCB", 1])

    dr = wb.create_sheet("DRG")
    dr.append(["标题"])
    dr.append(["DRG", "DRG名称", "DRG规则", "所属ADRG编码", "所属MDC编码", "排序"])
    dr.append(["AA19", "心肺移植", "", "AA1", "MDCA", 1])
    dr.append(["BB19", "兜底", "QTZD in MCC", "BB1", "MDCB", 1])

    st = wb.create_sheet("集合")
    st.append(["标题"])
    st.append(["集合编号", "ICD编码", "ICD名称", "类型编号"])
    st.append(["OP_AA1", "33.6x00", "心脏-肺联合移植术", "OP"])
    st.append(["DI_B00", "A01.001", "示例", "DI"])

    cc = wb.create_sheet("CC")
    cc.append(["标题"])
    cc.append(["疾病编码", "疾病名称", "排除表", "类型"])
    cc.append(["A01.001", "示例", "EX_1", "MCC"])

    ex = wb.create_sheet("排除表")
    ex.append(["标题"])
    ex.append(["集合编号", "ICD编码", "ICD名称"])
    ex.append(["EX_1", "A01.001", "示例"])

    wb.save(path)


def test_importer_round_trip(tmp_path: Path):
    init_db()
    xlsx = tmp_path / "mini.xlsx"
    _make_minimal_xlsx(xlsx)

    session = SessionLocal()
    try:
        summary = import_rules_xlsx(xlsx, session)

        assert summary["mdc"] == 2
        assert summary["adrg"] == 2
        assert summary["drg"] == 2
        assert summary["sets"] == 2
        assert summary["cc"] == 1
        assert summary["exclusion"] == 1
        assert summary["parse_errors"] == []

        # AST 缓存验证
        aa1 = session.query(AdrgRule).filter_by(adrg_code="AA1").one()
        ast = json.loads(aa1.parsed_ast)
        assert ast["type"] == "In"
        assert ast["variables"] == ["ZYSS"]
        assert ast["set_id"] == "OP_AA1"

        # 兜底 '1' → ConstTrueNode
        bb1 = session.query(AdrgRule).filter_by(adrg_code="BB1").one()
        ast = json.loads(bb1.parsed_ast)
        assert ast["type"] == "True"

        # 空规则的 DRG 也归一为 True
        aa19 = session.query(DrgRule).filter_by(drg_code="AA19").one()
        ast = json.loads(aa19.parsed_ast)
        assert ast["type"] == "True"

        # QTZD in MCC
        bb19 = session.query(DrgRule).filter_by(drg_code="BB19").one()
        ast = json.loads(bb19.parsed_ast)
        assert ast["type"] == "In"
        assert ast["set_id"] == "MCC"

        # 集合成员验证
        sm = session.query(SetMember).filter_by(set_id="OP_AA1").one()
        assert sm.code == "33.6X00"  # 已被 .upper() 归一
    finally:
        session.close()


def test_importer_handles_parse_errors_gracefully(tmp_path: Path):
    """parse 失败时不抛异常，只是记录。"""
    init_db()
    xlsx = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bad = wb.create_sheet("ADRG")
    bad.append(["标题"])
    bad.append(["ADRG编码", "ADRG名称", "ADRG规则", "所属MDC编码", "MDC内排序"])
    bad.append(["XX1", "bad", "ZYZD @@@ DI_B00", "MDCA", 1])
    wb.save(xlsx)

    # 同时要有其他表，否则报「期望 MDC 表存在」之类的
    for name, rows in [
        ("MDC", [["标题"], ["MDC编码", "MDC名称", "MDC规则", "排序"],
                 ["MDCA", "X", "", 1]]),
        ("DRG", [["标题"], ["DRG", "DRG名称", "DRG规则", "所属ADRG编码", "所属MDC编码", "排序"],
                 ["XX19", "X", "", "XX1", "MDCA", 1]]),
        ("集合", [["标题"], ["集合编号", "ICD编码", "ICD名称", "类型编号"]]),
        ("CC", [["标题"], ["疾病编码", "疾病名称", "排除表", "类型"]]),
        ("排除表", [["标题"], ["集合编号", "ICD编码", "ICD名称"]]),
    ]:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(xlsx)

    session = SessionLocal()
    try:
        summary = import_rules_xlsx(xlsx, session)
        assert len(summary["parse_errors"]) == 1
        assert summary["parse_errors"][0]["code"] == "XX1"
    finally:
        session.close()
