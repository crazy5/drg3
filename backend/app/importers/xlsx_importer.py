"""xlsx → SQLite 导入器。

读《DRG 3.0 分组方案配置信息.xlsx》的 6 张表，写入 SQLite，
并把所有规则 parse 为 AST 缓存到 parsed_ast 列。

idempotent：调用多次等价于一次（先清后写）。
"""
from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy.orm import Session

from app.db.models import (
    AdrgRule,
    CcList,
    DrgRule,
    ExclusionTable,
    MdcRule,
    SetMember,
)
from app.engine.parser.ast_nodes import serialize
from app.engine.parser.parser import parse_rule

log = logging.getLogger(__name__)

# 预 MDC：先期分组 / 新生儿 / 多发创伤
PRE_MDCS = {"MDCA", "MDCP", "MDCZ"}


def _norm_code(s: Any) -> str:
    return str(s).strip().upper() if s is not None else ""


def import_rules_xlsx(xlsx_path: str | Path, session: Session) -> dict:
    """导入整个 xlsx，返回统计摘要。"""
    p = Path(xlsx_path)
    if not p.exists():
        raise FileNotFoundError(p)

    wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    summary: dict[str, Any] = {
        "mdc": 0, "adrg": 0, "drg": 0,
        "sets": 0, "cc": 0, "exclusion": 0,
        "parse_errors": [],
    }

    # ─── MDC ───
    session.query(MdcRule).delete()
    for row in wb["MDC"].iter_rows(min_row=3, values_only=True):
        # (MDC编码, MDC名称, MDC规则, 排序)
        code, name, expr, priority = row[0], row[1], row[2], row[3]
        if not code:
            continue
        ast_blob, err = _safe_parse(expr)
        if err:
            summary["parse_errors"].append({"table": "MDC", "code": code, "error": err})
        session.add(MdcRule(
            mdc_code=_norm_code(code),
            mdc_name=name or "",
            rule_expr=expr or "",
            priority=int(priority or 0),
            is_pre_mdc=_norm_code(code) in PRE_MDCS,
            parsed_ast=ast_blob,
        ))
        summary["mdc"] += 1

    # ─── ADRG ───
    session.query(AdrgRule).delete()
    for row in wb["ADRG"].iter_rows(min_row=3, values_only=True):
        # (ADRG编码, ADRG名称, ADRG规则, 所属MDC编码, MDC内排序, ..., ...)
        code, name, expr, mdc, priority = row[0], row[1], row[2], row[3], row[4]
        if not code:
            continue
        ast_blob, err = _safe_parse(expr)
        if err:
            summary["parse_errors"].append({"table": "ADRG", "code": code, "error": err})
        adrg_type = _infer_adrg_type(expr or "")
        session.add(AdrgRule(
            adrg_code=_norm_code(code),
            adrg_name=name or "",
            mdc_code=_norm_code(mdc),
            rule_expr=expr or "",
            priority=int(priority or 0),
            adrg_type=adrg_type,
            parsed_ast=ast_blob,
        ))
        summary["adrg"] += 1

    # ─── DRG ───
    session.query(DrgRule).delete()
    for row in wb["DRG"].iter_rows(min_row=3, values_only=True):
        # (DRG, DRG名称, DRG规则, 所属ADRG编码, 所属MDC编码, 排序)
        code, name, expr, adrg, mdc, priority = (
            row[0], row[1], row[2], row[3], row[4], row[5],
        )
        if not code:
            continue
        # 空规则（继承 ADRG 名）→ AST 为 ConstTrueNode，但语义上保留一份空 expr
        ast_blob, err = _safe_parse(expr)
        if err:
            summary["parse_errors"].append({"table": "DRG", "code": code, "error": err})
        session.add(DrgRule(
            drg_code=_norm_code(code),
            drg_name=name or "",
            adrg_code=_norm_code(adrg),
            rule_expr=expr or "",
            priority=int(priority or 0),
            parsed_ast=ast_blob,
        ))
        summary["drg"] += 1

    # ─── 集合 ───
    session.query(SetMember).delete()
    seen: set[tuple[str, str]] = set()
    for row in wb["集合"].iter_rows(min_row=3, values_only=True):
        # (集合编号, ICD编码, ICD名称, 类型编号)
        set_id, code, name, type_no = row[0], row[1], row[2], row[3]
        if not set_id or not code:
            continue
        key = (_norm_code(set_id), _norm_code(code))
        if key in seen:
            continue
        seen.add(key)
        session.add(SetMember(
            set_id=key[0],
            code=key[1],
            code_name=name or "",
            code_type=type_no or "",
        ))
        summary["sets"] += 1

    # ─── CC ───
    session.query(CcList).delete()
    cc_seen: set[tuple[str, str]] = set()
    for row in wb["CC"].iter_rows(min_row=3, values_only=True):
        # (疾病编码, 疾病名称, 排除表, 类型)
        code, name, excl, level = row[0], row[1], row[2], row[3]
        if not code or not level:
            continue
        key = (_norm_code(code), (level or "").strip().upper())
        if key in cc_seen:
            continue
        cc_seen.add(key)
        session.add(CcList(
            diagnosis_code=key[0],
            level=key[1],
            exclusion_table=_norm_code(excl),
        ))
        summary["cc"] += 1

    # ─── 排除表 ───
    session.query(ExclusionTable).delete()
    excl_seen: set[tuple[str, str]] = set()
    for row in wb["排除表"].iter_rows(min_row=3, values_only=True):
        # (集合编号, ICD编码, ICD名称)
        table_id, code, name = row[0], row[1], row[2]
        if not table_id or not code:
            continue
        key = (_norm_code(table_id), _norm_code(code))
        if key in excl_seen:
            continue
        excl_seen.add(key)
        session.add(ExclusionTable(table_id=key[0], code=key[1]))
        summary["exclusion"] += 1

    session.commit()

    log.info(
        "imported: mdc=%d adrg=%d drg=%d sets=%d cc=%d exclusion=%d parse_errors=%d",
        summary["mdc"], summary["adrg"], summary["drg"],
        summary["sets"], summary["cc"], summary["exclusion"],
        len(summary["parse_errors"]),
    )
    return summary


def _safe_parse(expr: str | None) -> tuple[str, str | None]:
    """解析规则；返回 (AST JSON, error_or_None)。"""
    try:
        ast = parse_rule(expr)
        return json.dumps(serialize(ast), ensure_ascii=False), None
    except Exception as e:        # noqa: BLE001
        return json.dumps({"error": str(e), "raw": expr or ""}, ensure_ascii=False), str(e)


def _infer_adrg_type(expr: str) -> str:
    """根据规则粗略判断 ADRG 类型：surgical / medical / composite / fallback。"""
    s = expr.strip()
    if s in ("1", "1.0") or not s:
        return "fallback"
    if "ZYSS in" in s and "ZYZD in" in s:
        return "composite"
    if "ZYSS in" in s or "QTSS" in s:
        return "surgical"
    return "medical"
