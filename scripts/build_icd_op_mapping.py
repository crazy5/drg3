"""读 ICD9 国临版 3.0 对照医保版 2.0 的 xlsx，生成 icd_op_mapping.json。

xlsx 列：国临码 | 国临名 | 医保码 | 医保名

输出：
- backend/data/icd_op_mapping.json
  - key = 国临 3.0 手术码（归一：大写 + 去末尾点）
  - value = 医保 2.0 手术码
  - sync.py 拿这个值做 ICD9 国临→医保转换

- backend/data/icd_op_medical_name.json（二级兜底字典）
  - key = 医保码
  - value = 医保名
  - icd_dict.py 原码（医保码）未在主字典命中时来查这个，用于显示中文名

跟 build_icd_mapping.py 一致的 _norm()，确保 sync.py 能用同样的 key 查找。
"""
from __future__ import annotations

import json
import logging
import sys
from pathlib import Path

import openpyxl

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s: %(message)s")
log = logging.getLogger("icd_op_map")

REPO_ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = REPO_ROOT / "ICD9国临版3.0对照医保版2.0_0125.xlsx"
OUT_PATH = REPO_ROOT / "backend" / "data" / "icd_op_mapping.json"
OUT_MEDICAL = REPO_ROOT / "backend" / "data" / "icd_op_medical_name.json"


def _norm(code: str | None) -> str:
    """国临 3.0 码归一：大写 + 去末尾点。同步用作 sync.py 查询 key。"""
    if not code:
        return ""
    return str(code).strip().upper().rstrip(".")


def main() -> int:
    if not XLSX_PATH.exists():
        log.error("找不到对照表 xlsx: %s", XLSX_PATH)
        return 1

    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)
    if "Sheet1" not in wb.sheetnames:
        log.error("xlsx 里没找到 Sheet1（主对照表），实际 sheet: %s", wb.sheetnames)
        return 1
    ws = wb["Sheet1"]

    mapping: dict[str, str] = {}
    medical: dict[str, str] = {}        # 医保码 → 医保名（icd_dict 二级兜底）
    header_skipped = False
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 4:
            continue
        gl_code, _, yb_code, yb_name = row[0], row[1], row[2], row[3]
        # 跳过表头：表头第一列是「国临3.0手术代码」（字符串，不是数字）
        if not header_skipped:
            header_skipped = True
            if isinstance(gl_code, str) and not gl_code.replace(".", "").replace("X", "").isdigit():
                continue
        if not gl_code or not yb_code:
            continue
        k = _norm(gl_code)
        v = _norm(yb_code)
        if k and v:
            mapping[k] = v
        # 医保名（row[3]）：单独导出供 icd_dict 兜底
        if v and yb_name:
            name = str(yb_name).strip()
            if v not in medical and name:
                medical[v] = name

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(mapping, f, ensure_ascii=False, sort_keys=True)
    with open(OUT_MEDICAL, "w", encoding="utf-8") as f:
        json.dump(medical, f, ensure_ascii=False, sort_keys=True)

    log.info("生成 %s：%d 条映射", OUT_PATH, len(mapping))
    log.info("医保名二级表 %d 条 → %s", len(medical), OUT_MEDICAL)
    # 抽样
    sample = list(mapping.items())[:3]
    log.info("样例: %s", sample)
    # 看下 X 扩展（小写 x）保留情况
    x_ext_count = sum(1 for k in mapping if "X" in k)
    log.info("含 X 扩展的条目: %d", x_ext_count)
    return 0


if __name__ == "__main__":
    sys.exit(main())